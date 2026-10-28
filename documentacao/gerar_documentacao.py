from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
from datetime import datetime, timezone

OUT = Path(__file__).resolve().parent / "Sinal-documentacao-completa.docx"
doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)
styles = doc.styles
for name in ["Normal", "Body Text"]:
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(10.5)
    styles[name].font.color.rgb = RGBColor.from_string("31453B")
    styles[name].paragraph_format.space_after = Pt(7)
    styles[name].paragraph_format.line_spacing = 1.12
for name, size in [
    ("Title", 34),
    ("Heading 1", 23),
    ("Heading 2", 14),
    ("Heading 3", 11),
]:
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string("244B3B")
    styles[name].paragraph_format.space_before = Pt(12)
    styles[name].paragraph_format.space_after = Pt(7)
    styles[name].paragraph_format.keep_with_next = True
styles["Title"].paragraph_format.space_before = Pt(0)
styles["Subtitle"].font.name = "Calibri"
styles["Subtitle"].font.size = Pt(14)
styles["Subtitle"].font.color.rgb = RGBColor.from_string("748265")

# Brazilian Portuguese spelling language.
for style_name in [
    "Normal",
    "Title",
    "Subtitle",
    "Heading 1",
    "Heading 2",
    "Heading 3",
]:
    rpr = styles[style_name].element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "pt-BR")
    rpr.append(lang)
header = section.header.paragraphs[0]
r = header.add_run("sinal.")
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string("244B3B")
r = header.add_run("   |   GUIA DE FUNCIONAMENTO E DECISÕES TÉCNICAS")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string("748265")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Sinal • Reconhecimento ao vivo   |   ")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string("748265")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Sinal — como funciona e por que foi construído assim"
doc.core_properties.subject = (
    "Funcionamento, tecnologias, funções, arquitetura e decisões do projeto Sinal."
)
doc.core_properties.author = "Projeto Sinal"
doc.core_properties.keywords = (
    "Libras, React, TypeScript, MediaPipe, MySQL, SQL Server, Pandas, NumPy"
)
doc.core_properties.comments = ""
_now = datetime.now(timezone.utc)
doc.core_properties.created = _now
doc.core_properties.modified = _now


def p(text="", style=None):
    return doc.add_paragraph(text, style)


def h(text, level=2):
    return doc.add_heading(text, level)


def bullet(text):
    para = p(text, "List Bullet")
    para.paragraph_format.space_after = Pt(4)
    return para


def label(text):
    para = p(text)
    para.paragraph_format.space_after = Pt(8)
    run = para.runs[0]
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("9C6D50")


def box(title, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "EDF2E5")
    cell._tc.get_or_add_tcPr().append(shade)
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(7)
    para.paragraph_format.space_after = Pt(7)
    run = para.add_run(title + "\n")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("244B3B")
    para.add_run(text)
    row_no_split = OxmlElement("w:cantSplit")
    t.rows[0]._tr.get_or_add_trPr().append(row_no_split)
    # No trailing empty paragraph after a callout.


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
        sh = OxmlElement("w:shd")
        sh.set(qn("w:fill"), "244B3B")
        cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 0:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "F1F4EC")
                cell._tc.get_or_add_tcPr().append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(9)
        no_split = OxmlElement("w:cantSplit")
        t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    # Heading spacing separates tables from the next section.
    return t


def page(num, title):
    # Page break is attached to the next label, avoiding an empty overflow page.
    label(f"GUIA TÉCNICO  /  {num:02d}")
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    h(title, 1)


def code(text):
    para = p()
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(9)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F0F3EB")
    para._p.get_or_add_pPr().append(shade)
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


doc.core_properties.title = (
    "Sinal — documentação completa da versão com reconhecimento ao vivo"
)
doc.core_properties.subject = (
    "Funcionamento, funções, arquitetura, instalação, dados, modelo e limites."
)
label("DOCUMENTAÇÃO COMPLETA  /  VERSÃO COM RECONHECIMENTO AO VIVO")
h("Sinal\nDo gesto à letra", 0)
p("Guia do projeto, do código e das decisões técnicas.", "Subtitle")
p(
    "Este documento explica a versão que exibe “Letra X confirmada” sobre a câmera. A entrega inclui frontend React, modelo SpellNet convertido, serviços Python, esquemas de banco, testes, design system, ADRs e instruções de execução."
)
box(
    "A distinção mais importante",
    "O reconhecimento ao vivo atual usa imagens recortadas com o modelo pré-treinado SpellNet. O pipeline próprio de 63 coordenadas, com Pandas e NumPy, existe separadamente e ainda não substitui o SpellNet na câmera.",
)
h("Estado real da entrega")
table(
    ["Parte", "Situação"],
    [
        (
            "Câmera e reconhecimento",
            "Implementados: MediaPipe + SpellNet/ONNX no navegador, com confirmação temporal e limpeza ao perder a mão.",
        ),
        (
            "Tradução e prática",
            "Texto por adição de letras confirmadas; demonstração identificada; prática A–D com autoavaliação manual.",
        ),
        (
            "Serviços e bancos",
            "Código e Docker Compose preparados. MySQL/SQL Server não foram iniciados nem conectados ao front nesta entrega.",
        ),
        (
            "Modelo próprio",
            "Pipeline tabular experimental implementado; ainda depende de dados consentidos e avaliação antes de uso ao vivo.",
        ),
    ],
    [4.1, 12.9],
)
h("Como ler o guia")
p(
    "2. Execução e uso • 3. Webcam e MediaPipe • 4. Modelo e pré-processamento • 5. Confirmação temporal • 6. Arquitetura • 7. APIs e bancos • 8. Dados e treinamento • 9. Design system e UX • 10. Docker e conversão • 11. Testes e diagnóstico • 12. ADRs, glossário e fontes."
)
p(
    "A confirmação significa que a previsão permaneceu estável; não garante correção linguística. O Sinal explora o alfabeto manual e não traduz frases nem a gramática de Libras."
)
p("Referência do código: commit 47a6380e488536b42beb15757301cc95fa7f18f3.")

page(2, "Executar e usar o projeto")
h("O que vem no ZIP")
p(
    "A pasta Sinal contém src/, public/, services/, scripts/, contracts/, docs/, dist/, configurações e arquivos de dependências. O modelo ONNX está em public/models/spellnet. A pasta documentacao contém este Word e seu gerador."
)
p(
    "dist/ é o frontend já compilado. node_modules, ambiente virtual, histórico Git, caches e senhas reais não entram na entrega. As dependências podem ser reinstaladas pelos arquivos de configuração. O .env.example contém somente os campos a preencher."
)
h("Rodar o frontend para desenvolver")
p("Requer Node.js 24 LTS, npm e uma webcam. Dentro da pasta Sinal, execute:")
code("npm ci\nnpm run dev")
p(
    "Abra o endereço Local impresso pelo Vite. A câmera exige localhost ou HTTPS. Não abra index.html diretamente pelo explorador de arquivos."
)
h("Ver a versão compilada")
code("npm run build\nnpm run preview")
p(
    "O build já incluído também pode ser servido com Python, sem instalar as dependências do frontend:"
)
code("python -m http.server 4173 --directory dist")
p(
    "Depois, abra http://localhost:4173. Mesmo com o build local, o primeiro uso do detector MediaPipe precisa de acesso aos arquivos externos indicados na página 3."
)
h("Usar o reconhecimento")
bullet("Clique em Ativar câmera e autorize o acesso no navegador.")
bullet(
    "Posicione uma mão inteira, com boa iluminação, e mantenha o sinal por alguns instantes."
)
bullet(
    "Após a previsão estabilizar, veja “Letra X confirmada” no vídeo e a letra no painel."
)
bullet(
    "Clique em Adicionar letra ou use Enter para compor o texto; copie ou baixe o resultado."
)
bullet(
    "Retire a mão ou mude o gesto: a confirmação anterior deve desaparecer até uma nova estabilização."
)
p(
    "O modo prática continua manual: “Consegui fazer o sinal” registra sua própria avaliação. A demonstração usa exemplos ilustrativos e não mede a imagem da câmera."
)

page(3, "Webcam e detecção de pontos")
h("Captura nativa do navegador")
p(
    "navigator.mediaDevices.getUserMedia() solicita apenas vídeo. O pedido acontece depois do clique, não ao abrir a página. browserCamera implementa a integração, mantendo a API do navegador fora do domínio."
)
table(
    ["Função / estrutura", "Responsabilidade e escolha"],
    [
        (
            "useCamera()",
            "Controla câmera, vídeo, detector e estados. Centraliza a limpeza dos recursos ao cancelar, desligar ou sair da área da câmera.",
        ),
        (
            "start()",
            "Solicita MediaStream, conecta ao vídeo e aguarda play(). A geração da solicitação impede que uma autorização tardia reative uma captura cancelada.",
        ),
        (
            "stop()",
            "Encerra todas as tracks, limpa srcObject e os pontos, e invalida solicitações anteriores.",
        ),
        (
            "cameraError()",
            "Traduz falhas em mensagens: acesso negado, webcam ausente, ocupada ou navegador sem suporte adequado.",
        ),
        (
            "createHandDetector()",
            "Importa MediaPipe sob demanda e cria um HandLandmarker em modo VIDEO, configurado para uma mão.",
        ),
    ],
    [4.2, 12.8],
)
h("O que são os landmarks")
p(
    "O detector pré-treinado encontra 21 pontos: punho e articulações dos dedos. Cada um tem x, y e z; uma mão produz 63 números. MediaPipe localiza a mão; ele não é o classificador de letras do projeto."
)
code(
    "type Landmark = { x: number; y: number; z: number };\n21 pontos × 3 coordenadas = 63 valores"
)
h("Cadência, vídeo e espelhamento")
p(
    "requestAnimationFrame() acompanha a exibição do vídeo. A aplicação só analisa um frame novo e respeita um intervalo mínimo de cerca de 80 ms entre detecções. Isso limita o trabalho sem prometer FPS fixo em todos os aparelhos."
)
p(
    "A sobreposição usa HAND_CONNECTIONS para desenhar as ligações entre pontos. O espelhamento é visual: o modelo recebe a imagem original. Desligar “Mostrar pontos” oculta o desenho da tela, mas não desativa os pontos usados na classificação."
)
h("Arquivos carregados no primeiro uso")
p(
    "O detector baixa WASM de cdn.jsdelivr.net e o modelo hand_landmarker.task de storage.googleapis.com. A imagem da webcam permanece local. Se o carregamento falhar, o vídeo pode continuar disponível, mas o reconhecimento depende da detecção da mão."
)

page(4, "Como o modelo reconhece uma letra")
h("Modelo atual: SpellNet Libras")
p(
    "O modelo fornecido é uma CNN baseada em MobileNet, publicada por RAIA / SignLink em formato Keras. Para executá-lo sem servidor, o projeto o converteu para ONNX e usa ONNX Runtime Web com WebAssembly."
)
p(
    "As classes seguem a ordem do código original, não o alfabeto completo por suposição:"
)
code("A B C D E F G I L M N O P Q R S T U V W Y")
p(
    "São 21 classes estáticas. H, J, K, X e Z não estão nesse modelo. A biblioteca recebe os scores de saída nessa ordem; trocar a ordem poderia exibir uma letra errada mesmo com o modelo funcionando."
)
h("prepareFrame(): preparar exatamente a entrada esperada")
table(
    ["Etapa", "O que acontece"],
    [
        (
            "1. Validar",
            "Confere dimensões do vídeo e a presença dos 21 pontos. Não classifica sem mão detectada.",
        ),
        (
            "2. Recortar",
            "Calcula os limites da mão em pixels, acrescenta 100 pixels de margem e respeita os limites da imagem.",
        ),
        ("3. Redimensionar", "Desenha o recorte num Canvas de 224 × 224 pixels."),
        (
            "4. Marcar pontos",
            "Converte as posições para o recorte; desenha conexões verdes e pontos vermelhos, seguindo o processamento publicado.",
        ),
        (
            "5. Ajustar canais",
            "Lê RGBA do Canvas, usa a ordem BGR do fluxo OpenCV original e descarta o canal alfa.",
        ),
        (
            "6. Normalizar pixels",
            "Aplica pixel / 127,5 − 1; os valores ficam aproximadamente entre −1 e 1.",
        ),
    ],
    [3.5, 13.5],
)
h("createSpellNetClassifier() e selectPrediction()")
p(
    "O adaptador cria uma sessão ONNX, monta o tensor float32 com formato [1, 224, 224, 3] e executa predict(). selectPrediction() verifica o tamanho e os valores da saída e encontra o maior score. O resultado contém letter, confidence e modelVersion."
)
p(
    "O runtime usa uma thread WASM e os arquivos são empacotados pelo Vite. A sessão e os tensores são liberados ao encerrar o reconhecimento. O modelo tem aproximadamente 17 MB; o runtime WASM, cerca de 12 MB."
)
box(
    "Duas normalizações diferentes",
    "No SpellNet, normalizam-se pixels da imagem. Na proposta de modelo próprio, normalizam-se as coordenadas da mão. São entradas diferentes e não podem ser trocadas sem retreinar ou adaptar o modelo.",
)

page(5, "Quando aparece “Letra A confirmada”")
h("useRecognition(): coordenar a inferência")
p(
    "O hook carrega o classificador somente com a câmera ativa. Guarda a referência dos pontos mais recentes e realiza uma inferência de cada vez. Depois de concluir, agenda a próxima tentativa com intervalo de 220 ms; a velocidade real também depende do tempo gasto pelo modelo."
)
p(
    "Essa serialização evita acumular tarefas e classificar uma fila de frames antigos. Ao desligar, o hook cancela o agendamento, ignora respostas tardias e libera a sessão depois da inferência pendente."
)
h("RecognitionGate: filtrar oscilações")
p(
    "Um único frame pode produzir uma previsão passageira. A classe RecognitionGate mantém uma candidata e só libera a indicação quando todas as condições abaixo são atendidas:"
)
bullet("A classe pertence às 21 letras disponíveis.")
bullet("O score softmax é finito e pelo menos 0,80.")
bullet("Há no mínimo três observações consecutivas da mesma candidata.")
bullet("A candidata permanece a mesma por pelo menos 650 ms.")
code(
    "observe(previsao, agora)\n  → valida classe e score\n  → atualiza candidata, tempo e contagem\n  → retorna a letra somente após estabilidade"
)
p(
    "reset() apaga candidata, instante inicial e contagem. A aplicação faz isso ao perder a mão, cair abaixo do limiar ou interromper a câmera. Uma classe diferente reinicia a estabilização. Erros no detector também limpam os pontos para não manter uma letra antiga."
)
h("Como a tela usa o resultado")
p(
    "O App mostra a previsão estável sobre o vídeo e no painel “Sua tradução”. Enquanto o resultado ainda não estabilizou, exibe uma orientação. Uma letra confirmada pode ser adicionada ao texto por ação explícita; a aplicação não repete letras automaticamente a cada frame."
)
p(
    "Quando a mão deixa de ser detectada, a indicação é removida e o painel volta ao estado de espera. A demonstração mantém exemplos claramente identificados, separados do resultado real."
)
box(
    "O significado da confirmação",
    "O limiar 0,80 é uma regra de interface aplicada ao score do modelo. Não significa 80% de precisão comprovada. “Confirmada” indica concordância temporal das previsões, não certificação de que o sinal foi feito corretamente.",
)

page(6, "Arquitetura limpa e organização")
h("Separar regras, integrações e telas")
table(
    ["Pasta", "O que contém"],
    [
        (
            "src/domain/",
            "Landmark, Prediction, normalização, portas, classes do modelo e filtro temporal. Regras sem dependência de React ou MediaPipe.",
        ),
        (
            "src/application/",
            "Hooks de câmera, reconhecimento, prática, tradução e armazenamento, com dependências injetadas.",
        ),
        (
            "src/infrastructure/",
            "browserCamera, detector MediaPipe e adaptador SpellNet/ONNX. APIs de navegador e bibliotecas externas ficam aqui.",
        ),
        (
            "src/presentation/",
            "App organiza páginas e painéis. Componentes e estilos separados por responsabilidade; useWorkspace compõe os estados visuais.",
        ),
        (
            "services/",
            "Dois serviços Python independentes, cada um com domínio ou pipeline, API e repositório próprios.",
        ),
        (
            "docs/ e contracts/",
            "Arquitetura, design system, verificação, glossário, decisões e contratos HTTP/CSV.",
        ),
    ],
    [4, 13],
)
h("Portas e adaptadores")
p(
    "Uma porta descreve um contrato. Um adaptador implementa a integração concreta. Isso reduz o acoplamento: o domínio não precisa conhecer os detalhes de uma webcam ou de uma biblioteca de inferência."
)
table(
    ["Porta", "Contrato"],
    [
        ("CameraPort", "start() devolve a captura; stop() encerra a captura."),
        (
            "HandDetectorPort",
            "detect() devolve pontos ou null; dispose() libera o detector.",
        ),
        (
            "ClassifierPort",
            "Contrato reservado para o modelo próprio: recebe coordenadas normalizadas e devolve Prediction.",
        ),
        (
            "FrameClassifierPort",
            "Contrato do SpellNet atual: recebe vídeo e landmarks para o recorte; devolve Prediction; dispose() libera a sessão.",
        ),
    ],
    [4, 13],
)
h("Por que dois contratos de classificação?")
p(
    "O modelo original do Sinal foi pensado para 63 coordenadas. O SpellNet selecionado usa pixels. Criar FrameClassifierPort preserva essa diferença em vez de fingir que ambos recebem os mesmos dados."
)
p(
    "O frontend permanece uma aplicação React. A orientação a microserviços aparece nos limites dos serviços Python, na propriedade dos bancos e nos contratos versionados. Não foi adicionado um broker ou Kubernetes sem uma necessidade atual."
)
p(
    "main.tsx injeta AppServices criado em bootstrap/services.ts. Os hooks dependem de contratos, incluindo StoragePort e TextPort. npm run check verifica automaticamente os limites entre camadas."
)

page(7, "Serviços, MySQL e SQL Server")
h("Prática: MySQL")
p(
    "O serviço practice usa FastAPI e MySQL 8.4. PracticeSession.create() valida a letra A–D e tentativas de 1 a 100, gera um UUID e define evaluation como manual. MySQLPracticeRepository salva com parâmetros SQL e lista o catálogo."
)
table(
    ["Tabela", "Dados"],
    [
        (
            "practice_letters",
            "Letra e instrução. A chave da letra é referenciada pelas sessões.",
        ),
        (
            "practice_sessions",
            "UUID, letra, tentativas, avaliação manual e data. Há chave estrangeira e restrições para avaliação e tentativas.",
        ),
    ],
    [4.5, 12.5],
)
h("Treinamento: SQL Server")
p(
    "O serviço training usa FastAPI e SQL Server 2022 Developer. SQLServerExperimentRepository registra o resumo do dataset e o experimento numa transação. Os JSONs armazenam o baseline próprio e suas métricas, não as imagens da câmera."
)
table(
    ["Tabela", "Dados"],
    [
        (
            "dataset_versions",
            "SHA-256 do CSV, quantidade de linhas, quantidade de participantes e data.",
        ),
        (
            "experiments",
            "UUID, referência ao dataset, parâmetros em JSON, métricas em JSON e data.",
        ),
    ],
    [4.5, 12.5],
)
h("Contratos HTTP")
table(
    ["Serviço", "Rotas"],
    [
        (
            "Prática • porta 8001",
            "GET /v1/letters\nPOST /v1/sessions com letter e attempts",
        ),
        (
            "Treinamento • porta 8002",
            "POST /v1/experiments com dataset_csv e seed\nGET /v1/experiments/{uuid}",
        ),
        (
            "Ambos",
            "GET /health: processo ativo\nGET /ready: conexão e esquema do banco\n/docs e /openapi.json: documentação dos contratos",
        ),
    ],
    [4.5, 12.5],
)
p(
    "FastAPI valida entradas com Pydantic. SQLAlchemy gerencia conexões e consultas, usando PyMySQL para MySQL e pymssql para SQL Server. Dados inválidos retornam 422; indisponibilidade de banco retorna 503; experimento ausente retorna 404."
)
h("Por que essa divisão?")
p(
    "MySQL e SQL Server atendem ao requisito de stack, com responsabilidades próprias para evitar duplicação obrigatória. Cada serviço é dono de seu banco; não faz consultas cruzadas. A escolha permite evolução independente, mas aumenta o custo de operação."
)
box(
    "Estado da integração",
    "As APIs são uma base local. O front usa localStorage e não chama esses serviços. Eles não são necessários para reconhecer letras; autenticação e integração consentida ainda precisam ser implementadas antes de exposição externa.",
)

page(8, "Pandas, NumPy e o modelo próprio")
h("O formato do dataset")
p(
    "O pipeline recebe CSV de coordenadas. O arquivo deve ter exatamente letter, participant_id e x0,y0,z0 até x20,y20,z20. participant_id é um pseudônimo de até 64 caracteres em letras, números, sublinhado e hífen."
)
p(
    "A fase inicial aceita A–D, pelo menos duas classes e três participantes. Treino e teste precisam conter as mesmas letras; cada letra precisa de pelo menos dez exemplos no conjunto de treino."
)
h("train_baseline(): responsabilidades")
table(
    ["Etapa", "Funções e operações"],
    [
        (
            "Leitura e validação",
            "pd.read_csv() lê o CSV; isin(), isna(), fullmatch() e pd.to_numeric() verificam schema, rótulos e valores. Colunas extras e dados inválidos são rejeitados.",
        ),
        (
            "Normalização",
            "normalize_landmarks() usa arrays NumPy e operações vetoriais para centralizar e ajustar os lotes de mãos.",
        ),
        (
            "Divisão",
            "np.random.default_rng(seed) embaralha participantes. O teste recebe max(1, participantes // 4) participantes completos; o restante fica no treino.",
        ),
        (
            "Modelo baseline",
            "Calcula o centroide de cada classe pela média. Distâncias euclidianas ao centroide determinam a classe prevista.",
        ),
        (
            "Saída",
            "Exporta modelo, métricas e SHA-256 do CSV. O baseline é marcado validatedForProduction=false.",
        ),
    ],
    [3.6, 13.4],
)
h("normalizeLandmarks() e normalize_landmarks()")
p(
    "TypeScript e Python aplicam a mesma regra wrist-max-radius-v1: subtrair o punho de cada ponto e dividir todos pela maior distância 3D até o punho. O resultado é um vetor em ordem x0,y0,z0,...,x20,y20,z20."
)
code("qᵢ = pᵢ − p₀\ns = maior distância entre qᵢ e a origem\nvetor = qᵢ / s")
p(
    "A função TypeScript usa map(), Math.hypot(), Math.max() e flatMap(). Python usa broadcasting e np.linalg.norm(). Ambas rejeitam quantidade ou formato incorreto, coordenadas não finitas e mãos degeneradas."
)
p(
    "A regra remove translação e escala uniforme; não resolve rotação, lateralidade ou movimento. Separar os participantes evita avaliar o modelo com a mesma pessoa que ele já viu no treino."
)
h("Métricas e limites")
p(
    "O pipeline calcula accuracy, macroRecall e recallByLetter, além das contagens do split. O baseline de centroides é uma referência experimental, sem confiança calibrada. O JSON exportado não é instalado no frontend automaticamente e não altera o SpellNet ao vivo."
)

page(9, "Design system, UX e privacidade")
h("Identidade e componentes")
p(
    "O design usa verde profundo, fundo claro, tons vegetais e um detalhe terracota. DM Sans organiza a interface; Instrument Serif dá expressão aos títulos. As fontes são locais, com licenças OFL. Lucide fornece os ícones."
)
table(
    ["Elemento", "Função"],
    [
        (
            "Tokens CSS",
            "Centralizam cores, espaço, tipografia e raios em styles.css. Permitem ajustes consistentes.",
        ),
        (
            "Button",
            "Variantes primary, secondary e ghost, ícone opcional e disabled nativo.",
        ),
        ("Tag", "Estados neutral, green e orange com legenda textual."),
        ("Toggle", "Checkbox associado a rótulo, descrição e foco visível."),
        (
            "HandIllustration",
            "Referências esquemáticas A–D. Não equivalem a material linguístico validado.",
        ),
    ],
    [3.8, 13.2],
)
h("Fluxos da interface")
p(
    "A câmera tem estados explícitos de espera, solicitação, captura e erro. O reconhecimento tem idle, loading, ready e error. Há orientação enquanto o modelo carrega ou o gesto não estabiliza, em vez de mostrar uma letra arbitrária."
)
p(
    "A interface contém demonstração, prática manual, alfabeto inicial A–D, progresso, preferências e página do projeto. O diálogo de ajuda usa dialog nativo; as abas aceitam teclado; o menu mobile é ocultado quando fechado. O layout e as animações respeitam o espaço disponível e prefers-reduced-motion."
)
h("Persistência: useStoredState()")
p(
    "O hook tenta ler e validar o JSON do localStorage. Se o armazenamento não estiver disponível ou contiver dados inválidos, usa o valor inicial. As atualizações mantêm a interface utilizável mesmo se a gravação local falhar."
)
p(
    "As chaves sinal.preferences.v1 e sinal.practice.v1 guardam preferências e contagens de prática. Não existe sincronização entre aparelhos, conta de usuário ou garantia de preservação após limpar os dados do navegador."
)
h("Atalhos e texto")
p(
    "Com uma letra atual disponível, Enter adiciona a letra, espaço separa palavras e Backspace apaga, desde que o foco não esteja num campo ou botão. Os botões também permitem editar, copiar e baixar um arquivo de texto."
)
box(
    "Privacidade por construção",
    "Não é solicitado microfone. Vídeo e recortes são processados em memória no dispositivo e não são enviados às APIs. Os modelos precisam ser baixados. O endpoint de treinamento só recebe CSV quando alguém o envia explicitamente fora do fluxo atual do front.",
)

page(10, "Docker e reprodução dos modelos")
h("Subir os serviços opcionais")
p(
    "Requer Docker Compose. A imagem oficial de SQL Server usada é x86-64; reserve memória suficiente, no mínimo 2 GB só para o SQL Server. Developer é destinada a desenvolvimento e testes."
)
code(
    "# Linux / macOS\ncp .env.example .env\n\n# PowerShell\nCopy-Item .env.example .env"
)
p(
    "Preencha todas as senhas no .env. A senha administrativa do SQL Server deve atender à política do banco. A senha de aplicação usa 12–128 caracteres entre A–Z, a–z, 0–9, _ @ ! % -, com complexidade adequada. Não versione o arquivo preenchido."
)
code("docker compose --env-file .env up --build")
p(
    "O MySQL inicializa seu esquema no primeiro volume. O job sqlserver-init cria o login da aplicação, o banco e as tabelas. O serviço training não usa sa. As APIs ficam em 127.0.0.1:8001 e 127.0.0.1:8002; os bancos não publicam portas no host."
)
p(
    "Alterar senhas no .env não altera automaticamente os usuários dos volumes existentes. Faça a rotação no banco correspondente. Antes de exposição externa, ainda são necessários autenticação, autorização, limites, CORS explícito, backups e migrações."
)
h("Executar o pipeline Python sem banco")
code(
    "python -m venv .venv\n# Linux/macOS: source .venv/bin/activate\n# PowerShell: .\\.venv\\Scripts\\Activate.ps1\npip install -r requirements-dev.txt"
)
p("No Linux/macOS:")
code("PYTHONPATH=services/training python -m app dataset.csv --output baseline.json")
p(
    'No PowerShell, defina $env:PYTHONPATH="services/training" e execute python -m app dataset.csv --output baseline.json. O CSV não vem pronto: deve ser obtido de fonte autorizada ou coletado com consentimento.'
)
h("Reproduzir a conversão SpellNet → ONNX")
p(
    "A conversão já está entregue; não precisa ser repetida para usar a câmera. Para reproduzi-la, use um ambiente Python 3.12 separado, instale tensorflow-cpu==2.19.1, keras==3.8.0, tf2onnx==1.17.0 e onnxruntime==1.23.2 e obtenha o .keras original da fonte indicada na página 12."
)
code("python scripts/convert_spellnet.py modelo-original.keras public/models/spellnet")
p(
    "O script carrega com safe_mode=True e compile=False, converte com opset 17, compara saídas e argmax e gera o manifest.json com hashes. Isso verifica conversão matemática, não a qualidade de reconhecimento de Libras."
)

page(11, "Verificações e diagnóstico")
h("O que foi verificado")
table(
    ["Verificação", "Evidência e limite"],
    [
        (
            "Build e domínio frontend",
            "Tipos, arquitetura e build aprovados; 6 testes de domínio. Uma dependência proibida foi rejeitada na verificação de camadas.",
        ),
        (
            "Serviços Python",
            "18 testes de domínio, pipeline e HTTP. APIs recebem serviços e repositórios por injeção; verificam validação, falhas e readiness sem criar conexões reais.",
        ),
        (
            "Keras versus ONNX",
            "Duas entradas sintéticas, mesmo argmax e maior erro absoluto de aproximadamente 4,32 × 10⁻⁹. Não é métrica de precisão de sinais.",
        ),
        (
            "Fluxo real no navegador",
            "Stream de Canvas com imagens de referência A e B; MediaPipe e ONNX reais identificaram as letras. Verificados adição ao texto, limpeza sem mão e desligamento.",
        ),
        (
            "Experiência",
            "Na primeira entrega: demonstração, clipboard, download, persistência, menu mobile, teclado e ausência de overflow horizontal em 390/768/1024/1440 px.",
        ),
    ],
    [4.2, 12.8],
)
h("Como executar os testes")
code(
    "npm test\nnpm run build\n\n# Com o ambiente Python e requirements-dev.txt instalados:\nPYTHONPATH=services/practice python -m unittest discover -s services/practice/tests -v\nPYTHONPATH=services/training python -m unittest discover -s services/training/tests -v"
)
p(
    "No PowerShell, defina $env:PYTHONPATH para o serviço antes de cada comando python. Os roteiros de browser usados na entrega não foram adicionados como uma suíte portátil ao projeto; a evidência e os limites estão em docs/verification.md."
)
h("Problemas comuns")
table(
    ["Sintoma", "O que conferir"],
    [
        (
            "Câmera não abre",
            "Use localhost/HTTPS, libere a permissão do site e feche programas que estejam usando a webcam.",
        ),
        (
            "Reconhecimento carrega sem terminar",
            "Confira rede, download do modelo e arquivos WASM. Reative a câmera após corrigir a conexão.",
        ),
        (
            "Mão aparece, mas não confirma",
            "Melhore a luz, mostre a mão inteira, mantenha o gesto e confira se a letra pertence às 21 classes.",
        ),
        (
            "A letra está errada",
            "O modelo é experimental. Fundo, câmera, proporção, mão e estilo podem diferir dos dados originais. Não há precisão garantida.",
        ),
        (
            "API retorna 503",
            "Verifique o serviço e o esquema do banco. A interface local não depende desses endpoints.",
        ),
    ],
    [4.2, 12.8],
)
p(
    "A webcam física da pessoa usuária não foi exercitada neste ambiente. Não houve avaliação independente com novos sinalizantes, nem validação completa das 21 classes. O build estático não inclui bancos em execução."
)

page(12, "Decisões, glossário e referências")
h("ADRs entregues")
p(
    "001: processamento no navegador • 002: camadas e portas • 003: MySQL e SQL Server por serviço • 004: Pandas/NumPy e baseline de coordenadas • 005: design system • 006: reconhecimento ao vivo por SpellNet. 007: composição explícita, componentes e limites de dependência. O ADR 006 explica a entrada por imagem."
)
h("Glossário essencial")
table(
    ["Termo", "Significado"],
    [
        (
            "Libras / datilologia",
            "Língua Brasileira de Sinais / soletração por alfabeto manual, recorte deste projeto.",
        ),
        (
            "Landmark / feature",
            "Ponto estimado da mão / valor usado como entrada do modelo.",
        ),
        (
            "Inferência / treinamento",
            "Aplicar parâmetros já aprendidos / aprender parâmetros a partir de exemplos.",
        ),
        (
            "CNN / MobileNet",
            "Rede neural para padrões em imagens / arquitetura usada como base do SpellNet.",
        ),
        (
            "ONNX / WASM",
            "Formato de representação de modelos / tecnologia que executa o runtime no navegador.",
        ),
        (
            "Softmax / confirmação",
            "Scores relativos entre classes / concordância temporal exigida pela interface.",
        ),
        (
            "Holdout / centroide",
            "Conjunto separado para avaliação / vetor médio de uma classe.",
        ),
        (
            "Porta / adaptador",
            "Contrato de integração / implementação concreta desse contrato.",
        ),
        (
            "Microserviço / ADR",
            "Serviço com responsabilidade e dados próprios / registro de decisão arquitetural.",
        ),
    ],
    [4.2, 12.8],
)
h("Origem e atribuição")
p(
    "Modelo original: RAIA / SignLink. O Sinal não treinou o SpellNet. A versão consultada não apresentava licença explícita do modelo; essa ausência está registrada em NOTICE.txt e no ADR 006, sem atribuir uma licença presumida."
)
p(
    "Modelo e arquivo original:\nhttps://huggingface.co/RAIA-BRASIL/spellnet_libras\nCódigo e ordem das classes:\nhttps://github.com/gruporaia/SpellNet/blob/main/app/video_processing.py\nPré-processamento original:\nhttps://github.com/gruporaia/SpellNet/blob/main/app/preprocessing/signlink_preprocessing.py"
)
p(
    "O modelo convertido, os checksums, as classes e o erro de conversão estão em public/models/spellnet/manifest.json. A licença MIT do ONNX Runtime fica em public/runtime; as fontes locais incluem suas licenças OFL."
)
h("Evolução recomendada")
p(
    "Validar referências com pessoas surdas e profissionais de Libras; coletar coordenadas consentidas; avaliar generalização por participante; evoluir o classificador próprio; tratar letras com movimento por sequência temporal; e integrar os serviços somente após autenticação e consentimento definidos."
)

doc.save(OUT)
print("DOCX gerado:", OUT)
